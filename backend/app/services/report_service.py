import io
import os
import calendar
from datetime import datetime, date
from typing import Dict, Any, List, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as OpenpyxlImage

from app.db.mongodb import get_collection
from app.schemas.reports import AttendanceReportFilter

# Official Holiday Calendar 2026 matching b4S Solutions Pvt. Ltd. list
OFFICIAL_HOLIDAYS_2026 = {
    "2026-01-26": "Republic Day",
    "2026-03-04": "Dhulandi",
    "2026-03-19": "Gudi Padwa",
    "2026-05-01": "Maharashtra Day",
    "2026-08-15": "Independence Day",
    "2026-09-25": "Anant Chaturdashi",
    "2026-10-02": "Gandhi Jayanti",
    "2026-10-20": "Dussehra",
    "2026-11-08": "Diwali",
    "2026-11-09": "Goverdhan Pooja"
}

class ReportService:
    async def get_attendance_report(self, filters: AttendanceReportFilter) -> Dict[str, Any]:
        attendance_col = get_collection("attendance")
        conditions = []

        if filters.siteId and filters.siteId != "all":
            sites_col = get_collection("sites")
            site_doc = await sites_col.find_one({"$or": [{"id": filters.siteId}, {"_id": filters.siteId}, {"code": filters.siteId}]}) if sites_col is not None else None
            site_name = site_doc.get("name") if site_doc else None
            if site_name:
                conditions.append({"$or": [{"siteId": filters.siteId}, {"siteName": site_name}]})
            else:
                conditions.append({"siteId": filters.siteId})
        if filters.employeeId and filters.employeeId != "all":
            conditions.append({"employeeId": filters.employeeId})
        if filters.status and filters.status != "all":
            conditions.append({"status": filters.status})
        if filters.fromDate and filters.toDate:
            conditions.append({"date": {"$gte": filters.fromDate, "$lte": filters.toDate}})
        elif filters.fromDate:
            conditions.append({"date": {"$gte": filters.fromDate}})
        elif filters.toDate:
            conditions.append({"date": {"$lte": filters.toDate}})

        query: Dict[str, Any] = {"$and": conditions} if len(conditions) > 1 else (conditions[0] if conditions else {})

        records = await attendance_col.find(query).sort("date", -1).to_list(1000)
        formatted_records = []
        present_cnt = 0
        absent_cnt = 0
        late_cnt = 0
        half_day_cnt = 0
        checked_out_cnt = 0

        for r in records:
            item = dict(r)
            if "_id" in item:
                item["id"] = str(item["_id"])
                del item["_id"]
            status = item.get("status", "Present")
            if status == "Present":
                present_cnt += 1
            elif status == "Absent":
                absent_cnt += 1
            elif status == "Late":
                late_cnt += 1
            elif status == "Half Day":
                half_day_cnt += 1
            elif status == "Checked Out":
                checked_out_cnt += 1

            if not item.get("distanceTraveledFormatted") and item.get("checkInLocation") and item.get("checkOutLocation"):
                cin = item["checkInLocation"]
                cout = item["checkOutLocation"]
                if isinstance(cin, dict) and isinstance(cout, dict):
                    cin_lat, cin_lon = cin.get("latitude"), cin.get("longitude")
                    cout_lat, cout_lon = cout.get("latitude"), cout.get("longitude")
                    if cin_lat is not None and cout_lat is not None:
                        from app.utils.geo import calculate_haversine_distance, format_distance
                        d = calculate_haversine_distance(cin_lat, cin_lon, cout_lat, cout_lon)
                        item["distanceTraveledMeters"] = round(d, 1)
                        item["distanceTraveledFormatted"] = format_distance(d)

            formatted_records.append(item)

        # Enrich records with employee designation, circle, division
        emp_ids = list({r.get("employeeId") for r in formatted_records if r.get("employeeId")})
        if emp_ids:
            employees_col = get_collection("employees")
            emp_docs = await employees_col.find({"employeeId": {"$in": emp_ids}}).to_list(len(emp_ids))
            emp_map = {ed.get("employeeId"): ed for ed in emp_docs}
            for r in formatted_records:
                ed = emp_map.get(r.get("employeeId"))
                if ed:
                    r["designation"] = ed.get("designation") or "Field Technician"
                    r["circle"] = ed.get("circle") or "Prc"
                    r["division"] = ed.get("division") or ed.get("assignedSiteName") or "Office"
                else:
                    r["designation"] = "Field Technician"
                    r["circle"] = "Prc"
                    r["division"] = r.get("siteName") or "Office"

        total = len(formatted_records)
        compliance_pct = round((present_cnt + checked_out_cnt) / max(1, total) * 100, 1) if total > 0 else 0.0

        # Calculate real average hours from completed shifts if available
        completed_hours_total = 0
        completed_count = 0
        for r in formatted_records:
            wh = r.get("workingHours", "")
            if "h" in wh and "m" in wh:
                try:
                    parts = wh.replace("m", "").split("h")
                    h = int(parts[0].strip())
                    m = int(parts[1].strip())
                    completed_hours_total += (h * 60 + m)
                    completed_count += 1
                except Exception:
                    pass

        if completed_count > 0:
            avg_minutes = completed_hours_total // completed_count
            avg_hours_str = f"{avg_minutes // 60}h {avg_minutes % 60:02d}m"
        else:
            avg_hours_str = "0h 00m"

        return {
            "summary": {
                "totalRecords": total,
                "presentCount": present_cnt,
                "absentCount": absent_cnt,
                "lateCount": late_cnt,
                "halfDayCount": half_day_cnt,
                "checkedOutCount": checked_out_cnt,
                "averageWorkingHours": avg_hours_str,
                "complianceRate": f"{compliance_pct}%"
            },
            "records": formatted_records
        }

    async def generate_excel_bytes(self, filters: AttendanceReportFilter) -> bytes:
        """
        Generates an exact replica of the requested Adani Muster Roll attendance excel:
        - Columns: SR. NO., Employee Code, Employee Name, Designation, CIRCLE, Division, 1..31, P, WOP, A, OD, WO, PL, PH, Total
        - Title banner: ADANI ATTENDANCE FOR THE MONTH OF <MONTH> <YEAR>
        - Filled with real employee and attendance records from website database
        - Sheet 2: 'Detailed Log' with full telemetry timestamps, GPS, and employee division info
        """
        now = datetime.now()
        # Default to current month and year
        target_year = now.year
        target_month = now.month

        if getattr(filters, "year", None) and getattr(filters, "month", None):
            target_year = int(filters.year)
            target_month = int(filters.month)
        elif filters.toDate:
            try:
                parts = filters.toDate.split("-")
                target_year = int(parts[0])
                target_month = int(parts[1])
            except Exception:
                pass
        elif filters.fromDate:
            try:
                parts = filters.fromDate.split("-")
                target_year = int(parts[0])
                target_month = int(parts[1])
            except Exception:
                pass

        month_name = calendar.month_name[target_month]
        _, days_in_month = calendar.monthrange(target_year, target_month)

        # 1. Fetch Real Employees from website database (technicians & field staff)
        employees_col = get_collection("employees")
        emp_query: Dict[str, Any] = {"role": {"$ne": "admin"}}
        if filters.employeeId and filters.employeeId != "all":
            emp_query["employeeId"] = filters.employeeId
        if filters.siteId and filters.siteId != "all":
            sites_col = get_collection("sites")
            site_doc = await sites_col.find_one({"$or": [{"id": filters.siteId}, {"_id": filters.siteId}, {"code": filters.siteId}]}) if sites_col is not None else None
            site_target_name = site_doc.get("name") if site_doc else None
            match_or = [
                {"assignedSiteId": filters.siteId},
                {"assignedSiteIds": filters.siteId}
            ]
            if site_target_name:
                match_or.extend([
                    {"assignedSiteName": site_target_name},
                    {"assignedSiteNames": site_target_name}
                ])
            emp_query["$or"] = match_or

        employees = await employees_col.find(emp_query).sort("employeeId", 1).to_list(300)
        if not employees:
            # Fallback: if no employee strictly assigned to this site, include all active field technicians
            employees = await employees_col.find({"role": {"$ne": "admin"}}).sort("employeeId", 1).to_list(300)
        if not employees:
            # Fallback to users collection
            users_col = get_collection("users")
            user_query: Dict[str, Any] = {"role": {"$ne": "admin"}}
            if filters.employeeId and filters.employeeId != "all":
                user_query["employeeId"] = filters.employeeId
            employees = await users_col.find(user_query).sort("employeeId", 1).to_list(300)

        # 2. Fetch Sites map for location labels
        sites_col = get_collection("sites")
        sites_list = await sites_col.find({}).to_list(100)
        site_map = {}
        for s in sites_list:
            name = s.get("name")
            if s.get("id"):
                site_map[str(s["id"])] = name
            if s.get("_id"):
                site_map[str(s["_id"])] = name
            if s.get("code"):
                site_map[str(s["code"])] = name

        location_title = "All Locations"
        if filters.siteId and filters.siteId != "all":
            location_title = site_map.get(filters.siteId, filters.siteId)

        # 3. Fetch Real Attendance records for the target month
        attendance_col = get_collection("attendance")
        month_prefix = f"{target_year}-{target_month:02d}"
        att_query: Dict[str, Any] = {
            "date": {"$regex": f"^{month_prefix}"}
        }
        if filters.siteId and filters.siteId != "all":
            site_target_name = site_map.get(filters.siteId)
            if site_target_name:
                att_query["$or"] = [
                    {"siteId": filters.siteId},
                    {"siteName": site_target_name}
                ]
            else:
                att_query["siteId"] = filters.siteId
        if filters.employeeId and filters.employeeId != "all":
            att_query["employeeId"] = filters.employeeId

        att_records = await attendance_col.find(att_query).to_list(5000)

        # Build attendance lookup map: (employeeId, day_int) -> status_code
        att_map = {}
        for r in att_records:
            emp_id = r.get("employeeId")
            date_str = r.get("date", "")
            status_val = r.get("status", "Present")
            if emp_id and date_str:
                try:
                    d_int = int(date_str.split("-")[2])
                    code = "P"
                    if status_val in ["Present", "Checked Out"]:
                        code = "P"
                    elif status_val == "Absent":
                        code = "A"
                    elif status_val == "Late":
                        code = "L"
                    elif status_val == "Half Day":
                        code = "HD"
                    elif status_val in ["On Duty", "OD"]:
                        code = "OD"
                    elif status_val in ["Weekly Off", "WO", "O"]:
                        code = "O"
                    elif status_val in ["Holiday", "PH", "HL"]:
                        code = "HL"
                    elif status_val in ["Leave", "PL"]:
                        code = "PL"
                    elif status_val in ["WOP", "Without Pay"]:
                        code = "WOP"
                    att_map[(emp_id, d_int)] = code
                except Exception:
                    pass

        # 4. Create Excel Workbook with openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Attendance"
        ws.views.sheetView[0].showGridLines = True

        # Exact styling matching the muster roll template
        font_title = Font(name="Calibri", size=13, bold=True, color="1F497D")
        font_company = Font(name="Calibri", size=10, bold=True, color="595959")
        font_loc = Font(name="Calibri", size=10, bold=True, color="1F497D")
        font_header = Font(name="Calibri", size=9, bold=True, color="FFFFFF")
        font_body = Font(name="Calibri", size=9, bold=False, color="000000")
        font_body_bold = Font(name="Calibri", size=9, bold=True, color="000000")
        font_total = Font(name="Calibri", size=9, bold=True, color="1F497D")

        fill_header = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
        fill_totals_header = PatternFill(start_color="2A5C9A", end_color="2A5C9A", fill_type="solid")
        fill_sunday = PatternFill(start_color="F2F5F9", end_color="F2F5F9", fill_type="solid")
        fill_holiday = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
        fill_summary_col = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        fill_total_row = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")

        border_thin = Border(
            left=Side(style="thin", color="D9D9D9"),
            right=Side(style="thin", color="D9D9D9"),
            top=Side(style="thin", color="D9D9D9"),
            bottom=Side(style="thin", color="D9D9D9")
        )
        border_total = Border(
            left=Side(style="thin", color="B0B0B0"),
            right=Side(style="thin", color="B0B0B0"),
            top=Side(style="thin", color="000000"),
            bottom=Side(style="double", color="000000")
        )

        align_center = Alignment(horizontal="center", vertical="center")
        align_left = Alignment(horizontal="left", vertical="center")

        # 5. Row 1: Main Title
        ws.merge_cells("A1:AS1")
        cell_r1 = ws["A1"]
        cell_r1.value = f"ADANI ATTENDANCE FOR THE MONTH OF {month_name[:3].upper()} {target_year}"
        cell_r1.font = font_title
        cell_r1.alignment = align_left
        ws.row_dimensions[1].height = 24

        # Row 2: Subtitle
        ws.merge_cells("C2:AS2")
        cell_r2 = ws["C2"]
        cell_r2.value = f"                                 ATTENDANCE FOR THE MONTH OF {month_name.upper()}, {target_year}"
        cell_r2.font = font_company
        cell_r2.alignment = align_left
        ws.row_dimensions[2].height = 20

        # Row 3: Location
        cell_r3_lbl = ws["B3"]
        cell_r3_lbl.value = "LOCATION"
        cell_r3_lbl.font = font_loc
        cell_r3_lbl.alignment = align_left

        ws.merge_cells("C3:AS3")
        cell_r3_val = ws["C3"]
        cell_r3_val.value = location_title
        cell_r3_val.font = font_loc
        cell_r3_val.alignment = align_left
        ws.row_dimensions[3].height = 18

        # Row 4: Exact Requested Column Headers
        headers = [
            "SR. NO.",
            "Employee Code",
            "Employee Name",
            "Designation",
            "CIRCLE",
            "Division"
        ]
        # Days 1 to 31
        for d in range(1, 32):
            headers.append(d)
        # Summary counts
        headers.extend(["P", "WOP", "A", "OD", "WO", "PL", "PH", "Total"])

        ws.row_dimensions[4].height = 26
        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=4, column=col_idx, value=h)
            cell.font = font_header
            cell.alignment = align_center
            cell.border = border_thin
            if col_idx >= 38:
                cell.fill = fill_totals_header
            else:
                cell.fill = fill_header

        # Column Widths
        col_widths = {
            1: 8.0,    # SR. NO.
            2: 18.0,   # Employee Code
            3: 28.0,   # Employee Name
            4: 22.0,   # Designation
            5: 14.0,   # CIRCLE
            6: 22.0,   # Division
        }
        for d_col in range(7, 38):
            col_widths[d_col] = 4.5  # Days 1 to 31
        for sum_col in range(38, 45):
            col_widths[sum_col] = 5.5  # P, WOP, A, OD, WO, PL, PH
        col_widths[45] = 9.5  # Total

        for c_idx, width in col_widths.items():
            ws.column_dimensions[get_column_letter(c_idx)].width = width

        # Today date for real timeline mapping
        today_date = date.today()

        # Rows 5 onwards: Real Employee Data Rows
        start_row = 5
        for emp_idx, emp in enumerate(employees, 1):
            curr_row = start_row + emp_idx - 1
            ws.row_dimensions[curr_row].height = 20

            emp_code = emp.get("employeeId", f"b4S/Adani/{emp_idx:02d}")
            emp_name = emp.get("name") or emp.get("fullName", "Field Technician")
            emp_desig = emp.get("designation") or "Junior Engineer"

            emp_site = emp.get("assignedSiteName") or location_title
            emp_circle = emp.get("circle")
            if not emp_circle:
                if "Pune" in str(emp_site):
                    emp_circle = "Prc"
                elif "Mumbai" in str(emp_site) or "Bandra" in str(emp_site):
                    emp_circle = "Rastapeth"
                elif "Surat" in str(emp_site) or "Ahmedabad" in str(emp_site):
                    emp_circle = "Gujarat Circle"
                else:
                    emp_circle = "Prc"

            emp_division = emp.get("division")
            if not emp_division:
                if "Pune" in str(emp_site):
                    emp_division = "Pune Division"
                elif "Phase 1" in str(emp_site):
                    emp_division = "Office"
                elif "Phase 2" in str(emp_site):
                    emp_division = "Bundgarden"
                elif "Surat" in str(emp_site):
                    emp_division = "Surat Division"
                elif "Mumbai" in str(emp_site):
                    emp_division = "Office"
                elif emp_site and emp_site != "All Locations":
                    emp_division = emp_site
                else:
                    emp_division = "Office"

            # Col 1: SR. NO.
            c_sr = ws.cell(row=curr_row, column=1, value=emp_idx)
            c_sr.font = font_body
            c_sr.alignment = align_center
            c_sr.border = border_thin

            # Col 2: Employee Code
            c_code = ws.cell(row=curr_row, column=2, value=emp_code)
            c_code.font = font_body_bold
            c_code.alignment = align_left
            c_code.border = border_thin

            # Col 3: Employee Name
            c_name = ws.cell(row=curr_row, column=3, value=emp_name)
            c_name.font = font_body_bold
            c_name.alignment = align_left
            c_name.border = border_thin

            # Col 4: Designation
            c_desig = ws.cell(row=curr_row, column=4, value=emp_desig)
            c_desig.font = font_body
            c_desig.alignment = align_left
            c_desig.border = border_thin

            # Col 5: CIRCLE
            c_circle = ws.cell(row=curr_row, column=5, value=emp_circle)
            c_circle.font = font_body
            c_circle.alignment = align_left
            c_circle.border = border_thin

            # Col 6: Division
            c_div = ws.cell(row=curr_row, column=6, value=emp_division)
            c_div.font = font_body
            c_div.alignment = align_left
            c_div.border = border_thin

            # Days 1 to 31 (Cols 7 to 37)
            for day in range(1, 32):
                col_i = 6 + day
                cell_day = ws.cell(row=curr_row, column=col_i)
                cell_day.border = border_thin
                cell_day.alignment = align_center

                if day > days_in_month:
                    cell_day.value = ""
                    cell_day.font = font_body
                    continue

                curr_dt = date(target_year, target_month, day)
                dt_str = curr_dt.strftime("%Y-%m-%d")
                is_sunday = (curr_dt.weekday() == 6)
                is_holiday = dt_str in OFFICIAL_HOLIDAYS_2026

                code = att_map.get((emp_code, day))
                if not code:
                    if is_holiday:
                        code = "HL"
                    elif is_sunday:
                        code = "O"
                    elif curr_dt < today_date:
                        code = "A"
                    elif curr_dt == today_date:
                        code = "-"
                    else:
                        code = ""

                cell_day.value = code
                if code == "P":
                    cell_day.font = Font(name="Calibri", size=9, bold=True, color="047857")
                elif code == "A":
                    cell_day.font = Font(name="Calibri", size=9, bold=True, color="B91C1C")
                elif code in ["O", "WO"]:
                    cell_day.font = Font(name="Calibri", size=9, bold=True, color="1D4ED8")
                    cell_day.fill = fill_sunday
                elif code in ["HL", "PH"]:
                    cell_day.font = Font(name="Calibri", size=9, bold=True, color="B45309")
                    cell_day.fill = fill_holiday
                elif code == "OD":
                    cell_day.font = Font(name="Calibri", size=9, bold=True, color="6D28D9")
                elif code == "WOP":
                    cell_day.font = Font(name="Calibri", size=9, bold=True, color="047857")
                elif code == "PL":
                    cell_day.font = Font(name="Calibri", size=9, bold=True, color="D97706")
                else:
                    cell_day.font = font_body

            # Summary Columns (AL to AS / 38 to 45)
            # Days range: G{curr_row}:AK{curr_row}
            day_start_col = "G"
            day_end_col = "AK"

            formulas = {
                38: f'=COUNTIF({day_start_col}{curr_row}:{day_end_col}{curr_row},"P")',
                39: f'=COUNTIF({day_start_col}{curr_row}:{day_end_col}{curr_row},"WOP")+COUNTIF({day_start_col}{curr_row}:{day_end_col}{curr_row},"O")',
                40: f'=COUNTIF({day_start_col}{curr_row}:{day_end_col}{curr_row},"A")',
                41: f'=COUNTIF({day_start_col}{curr_row}:{day_end_col}{curr_row},"OD")',
                42: f'=COUNTIF({day_start_col}{curr_row}:{day_end_col}{curr_row},"WO")',
                43: f'=COUNTIF({day_start_col}{curr_row}:{day_end_col}{curr_row},"PL")',
                44: f'=COUNTIF({day_start_col}{curr_row}:{day_end_col}{curr_row},"PH")+COUNTIF({day_start_col}{curr_row}:{day_end_col}{curr_row},"HL")',
                45: f'=AL{curr_row}+AM{curr_row}+AO{curr_row}+AP{curr_row}+AQ{curr_row}+AR{curr_row}',
            }

            for col_s, formula in formulas.items():
                c_sum = ws.cell(row=curr_row, column=col_s, value=formula)
                c_sum.border = border_thin
                c_sum.alignment = align_center
                c_sum.font = font_body_bold
                c_sum.fill = fill_summary_col

        # Grand Total Row at Bottom
        num_emps = len(employees)
        tot_row = start_row + num_emps
        ws.row_dimensions[tot_row].height = 22

        c_tot_label = ws.cell(row=tot_row, column=1, value="Total")
        c_tot_label.font = font_total
        c_tot_label.alignment = align_center
        c_tot_label.border = border_total
        c_tot_label.fill = fill_total_row

        for c_blank in range(2, 38):
            c_b = ws.cell(row=tot_row, column=c_blank, value="")
            c_b.border = border_total
            c_b.fill = fill_total_row

        for sum_c in range(38, 46):
            col_let = get_column_letter(sum_c)
            if num_emps > 0:
                sum_form = f"=SUM({col_let}{start_row}:{col_let}{tot_row - 1})"
            else:
                sum_form = 0
            c_sum_tot = ws.cell(row=tot_row, column=sum_c, value=sum_form)
            c_sum_tot.font = font_total
            c_sum_tot.alignment = align_center
            c_sum_tot.border = border_total
            c_sum_tot.fill = fill_total_row

        # Sheet 2: Detailed Log Tab with full telemetry
        ws_detail = wb.create_sheet(title="Detailed Log")
        ws_detail.views.sheetView[0].showGridLines = True

        detail_headers = [
            "Date", "Employee ID", "Employee Name", "Designation", "Circle", "Division", "Site Name", "Status",
            "Check-In Time", "Check-Out Time", "Working Hours", "Shift Distance",
            "Check-In Latitude", "Check-In Longitude", "Check-Out Latitude", "Check-Out Longitude"
        ]

        ws_detail.row_dimensions[1].height = 24
        for col_idx, h in enumerate(detail_headers, 1):
            cell = ws_detail.cell(row=1, column=col_idx, value=h)
            cell.font = font_header
            cell.fill = fill_header
            cell.alignment = align_center
            cell.border = border_thin

        detail_widths = [14, 16, 24, 18, 14, 18, 22, 14, 15, 15, 15, 16, 18, 18, 18, 18]
        for idx, w in enumerate(detail_widths, 1):
            ws_detail.column_dimensions[get_column_letter(idx)].width = w

        emp_details_map = {e.get("employeeId"): e for e in employees if e.get("employeeId")}

        for r_idx, r in enumerate(att_records, 2):
            ws_detail.row_dimensions[r_idx].height = 19
            cin_loc = r.get("checkInLocation") or {}
            cout_loc = r.get("checkOutLocation") or {}

            # Distance formatted
            dist_fmt = r.get("distanceTraveledFormatted")
            if not dist_fmt and cin_loc and cout_loc:
                if isinstance(cin_loc, dict) and isinstance(cout_loc, dict):
                    c1_lat, c1_lon = cin_loc.get("latitude"), cin_loc.get("longitude")
                    c2_lat, c2_lon = cout_loc.get("latitude"), cout_loc.get("longitude")
                    if c1_lat is not None and c2_lat is not None:
                        from app.utils.geo import calculate_haversine_distance, format_distance
                        dist_fmt = format_distance(calculate_haversine_distance(c1_lat, c1_lon, c2_lat, c2_lon))

            e_info = emp_details_map.get(r.get("employeeId")) or {}
            e_desig = e_info.get("designation") or r.get("designation") or "Field Technician"
            e_circle = e_info.get("circle") or r.get("circle") or "Prc"
            e_div = e_info.get("division") or r.get("division") or r.get("siteName") or "Office"

            row_vals = [
                r.get("date", ""),
                r.get("employeeId", ""),
                r.get("employeeName", ""),
                e_desig,
                e_circle,
                e_div,
                r.get("siteName", ""),
                r.get("status", ""),
                r.get("checkInTime", "--"),
                r.get("checkOutTime", "--"),
                r.get("workingHours", "--"),
                dist_fmt or "--",
                cin_loc.get("latitude", "") if isinstance(cin_loc, dict) else "",
                cin_loc.get("longitude", "") if isinstance(cin_loc, dict) else "",
                cout_loc.get("latitude", "") if isinstance(cout_loc, dict) else "",
                cout_loc.get("longitude", "") if isinstance(cout_loc, dict) else ""
            ]

            for c_idx, val in enumerate(row_vals, 1):
                c = ws_detail.cell(row=r_idx, column=c_idx, value=val)
                c.font = font_body
                c.border = border_thin
                c.alignment = align_left if c_idx in [2, 3, 4, 5, 6, 7] else align_center

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output.getvalue()


report_service = ReportService()
